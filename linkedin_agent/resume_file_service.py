from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


GENERATED_RESUME_DIR = Path(
    "generated_resumes"
)

SECTION_HEADINGS = {
    "PROFESSIONAL SUMMARY",
    "SUMMARY",
    "EXPERIENCE",
    "PROFESSIONAL EXPERIENCE",
    "TECHNICAL SKILLS",
    "SKILLS",
    "EDUCATION",
    "CERTIFICATIONS",
    "AWARDS",
    "PROJECTS",
}


def ensure_resume_directory() -> Path:
    GENERATED_RESUME_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    return GENERATED_RESUME_DIR


def sanitize_filename(value: str) -> str:
    """
    Convert a company/title into a safe filename fragment.
    """

    value = value.lower().strip()

    value = re.sub(
        r"[^a-z0-9]+",
        "_",
        value,
    )

    return value.strip("_")

def add_section_heading(
    document: Document,
    text: str,
) -> None:
    """
    Add a clean ATS-friendly resume section heading.
    """

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(
        text.upper()
    )

    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_bullet(
    document: Document,
    text: str,
) -> None:
    """
    Add a resume achievement bullet.
    """

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.paragraph_format.left_indent = Inches(
        0.2
    )

    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        text
    )

    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def add_name_header(
    document: Document,
    name: str,
) -> None:

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        name
    )

    run.bold = True
    run.font.name = "Caliri"
    run.font.size = Pt(16)


def add_contact_line(
    document: Document,
    contact_line: str,
) -> None:

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(
        contact_line
    )

    run.font.name = "Calibri"
    run.font.size = Pt(9.5)



def add_body_paragraph(
    document: Document,
    text: str,
) -> None:

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(text)

    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def build_resume_filename(
    company: str,
    title: str,
    job_id: str,
    extension: str = "docx",
) -> Path:

    directory = ensure_resume_directory()

    company_name = sanitize_filename(
        company
    )

    title_name = sanitize_filename(
        title
    )

    filename = (
        f"{company_name}_"
        f"{title_name}_"
        f"{job_id}."
        f"{extension}"
    )

    return directory / filename


def add_job_header(
    document: Document,
    title: str,
    company: str,
    location: str | None = None,
    dates: str | None = None,
) -> None:
    """
    Add an ATS-friendly experience header.

    Example:
    Senior Engineering Manager | Walmart | Bentonville, AR | 2020–Present
    """

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(1)

    # Job title
    title_run = paragraph.add_run(title)

    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(10.5)

    # Company
    if company:
        company_run = paragraph.add_run(
            f" | {company}"
        )

        company_run.bold = True
        company_run.font.name = "Calibri"
        company_run.font.size = Pt(10.5)

    # Location
    if location:
        location_run = paragraph.add_run(
            f" | {location}"
        )

        location_run.font.name = "Calibri"
        location_run.font.size = Pt(9.5)

    # Dates
    if dates:
        dates_run = paragraph.add_run(
            f" | {dates}"
        )

        dates_run.italic = True
        dates_run.font.name = "Calibri"
        dates_run.font.size = Pt(9.5)


def looks_like_dates(
    value: str,
) -> bool:

    value_lower = value.lower()

    if "present" in value_lower:
        return True

    # Matches years such as 2019, 2024
    if re.search(
        r"\b(19|20)\d{2}\b",
        value
    ):
        return True

    return False


def parse_job_header(
    line: str,
) -> dict | None:
    """
    Parse:

    Title | Company | Location | Dates

    Returns None if the line does not look like a job header.
    """

    parts = [
        part.strip()
        for part in line.split("|")
    ]

    if len(parts) < 2:
        return None

    result = {
        "title": parts[0],
        "company": parts[1],
        "location": None,
        "dates": None,
    }

    if len(parts) == 3:
        # Could be location OR dates.
        # Check whether the third value looks like dates.
        if looks_like_dates(parts[2]):
            result["dates"] = parts[2]
        else:
            result["location"] = parts[2]

    elif len(parts) >= 4:
        result["location"] = parts[2]
        result["dates"] = parts[3]

    return result


def save_resume_text(
    resume_text: str,
    company: str,
    title: str,
    job_id: str,
) -> str:

    path = build_resume_filename(
        company=company,
        title=title,
        job_id=job_id,
        extension="txt",
    )

    path.write_text(
        resume_text,
        encoding="utf-8",
    )

    return str(
        path.resolve()
    )

def save_resume_docx(
    resume_text: str,
    company: str,
    title: str,
    job_id: str,
) -> str:

    path = build_resume_filename(
        company=company,
        title=title,
        job_id=job_id,
        extension="docx",
    )

    document = Document()

    # --------------------------------------------------
    # PAGE SETUP
    # --------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    # --------------------------------------------------
    # DEFAULT FONT
    # --------------------------------------------------

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    # --------------------------------------------------
    # CONTENT
    # --------------------------------------------------
    current_section = None

    lines = resume_text.splitlines()

    for raw_line in lines:

        line = raw_line.strip()

        if not line:
            continue

        # ==================================================
        # SECTION HEADINGS
        # ==================================================
        if line.upper() in SECTION_HEADINGS:
            
            current_section = line.upper()

            add_section_heading(
                document,
                line,
            )

            continue
       
    # ==================================================
    # EXPERIENCE HEADER
    # ==================================================

        if current_section in (
            "EXPERIENCE",
            "PROFESSIONAL EXPERIENCE",
        ):

            job_header = parse_job_header(
                line
            )

            if job_header:

                add_job_header(
                    document=document,
                    title=job_header["title"],
                    company=job_header["company"],
                    location=job_header["location"],
                    dates=job_header["dates"],
                )

                continue

        if line.startswith(
            ("•", "-", "*")
        ):

            clean_line = line.lstrip(
                "•-* "
            ).strip()

            add_bullet(
                document,
                clean_line,
            )

            continue

        add_body_paragraph(
            document,
            line,
        )

    document.save(path)

    return str(
        path.resolve()
    )